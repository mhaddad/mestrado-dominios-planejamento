"""Monta a dissertação revisada em .docx, no padrão do guia da FEI (redacao/README.md).

1. Pandoc junta os capítulos em Markdown (ordem em ARQUIVOS), resolve as citações com o referencias.bib e
   o CSL abnt-fei, aplica o filtro fei.lua e o modelo modelo-fei.docx.
2. O script acrescenta o que o Pandoc não faz:
   - capa, folha de rosto e a página da ficha catalográfica (identificacao.json);
   - campos do Word para o sumário e as listas de ilustrações, quadros e tabelas;
   - um espaço, e não tabulação, entre o número e o título das seções (guia, p. 4);
   - títulos sem número (REFERÊNCIAS, APÊNDICE) centralizados (guia, p. 4);
   - três seções de página: capa; pré-textuais contadas a partir da folha de rosto, sem número impresso;
     texto a partir da Introdução, com o número no canto superior direito em fonte 10 (guia, p. 5).
Os campos (sumário, listas, número de página) se atualizam ao abrir o arquivo no Word (F9, se pedir).

Uso: uv run --no-project --with python-docx python redacao/montagem/montar.py
Saída: redacao/saida/dissertacao-haddad-2026.docx
Gera a versão anterior à revisão do autor; a final é redacao/final/dissertacao-haddad-2026.docx (redacao/README.md).
"""
import copy
import json
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

RAIZ = Path(__file__).resolve().parents[2]
AQUI = RAIZ / "redacao/montagem"
CAPITULOS = RAIZ / "redacao/capitulos"
SAIDA = RAIZ / "redacao/saida/dissertacao-haddad-2026.docx"
ARQUIVOS = [
    "00-pretextuais.md",
    "01-introducao.md",
    "02-fundamentos.md",
    "03-revisitando-2010.md",
    "04-metodo.md",
    "05-resultados.md",
    "06-llms.md",
    "07-ponte-software.md",
    "08-conclusoes.md",
    "09-referencias.md",
    "10-apendices.md",
]
CAMPOS = {
    "sumario": r'TOC \o "1-5" \h \z \u',
    "ilustracoes": r'TOC \h \z \t "Legenda de figura,1"',
    "quadros": r'TOC \h \z \t "Legenda de quadro,1"',
    "tabelas": r'TOC \h \z \t "Table Caption,1"',
}


def pandoc(arquivos, destino):
    cmd = ["pandoc", *map(str, arquivos), "--from=markdown", "--citeproc",
           f"--bibliography={RAIZ / 'literatura/referencias/referencias.bib'}",
           f"--csl={RAIZ / 'redacao/estilos/abnt-fei.csl'}",
           f"--lua-filter={AQUI / 'fei.lua'}",
           f"--reference-doc={AQUI / 'modelo-fei.docx'}",
           f"--resource-path={CAPITULOS}",
           "--number-sections", "-o", str(destino)]
    subprocess.run(cmd, check=True)


def campo(paragrafo, instrucao, texto="Atualize o campo no Word (F9)."):
    """Troca o conteúdo do parágrafo por um campo do Word."""
    for r in list(paragrafo.runs):
        r._r.getparent().remove(r._r)
    def run_com(el):
        r = OxmlElement("w:r")
        r.append(el)
        paragrafo._p.append(r)
    f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "begin"); f.set(qn("w:dirty"), "true"); run_com(f)
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = f" {instrucao} "; run_com(i)
    f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "separate"); run_com(f)
    t = OxmlElement("w:t"); t.text = texto; run_com(t)
    f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "end"); run_com(f)


def espaco_no_numero(par):
    """O Pandoc separa número e título com tabulação; o guia pede um espaço."""
    for tab in list(par._p.iter(qn("w:tab"))):
        r = tab.getparent()
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = " "
        r.replace(tab, t)


def novo_paragrafo(ancora, texto="", negrito=False, alinhamento=WD_ALIGN_PARAGRAPH.CENTER, antes=0,
                   entrelinhas=1.5, recuo_esq=None, tamanho=12, quebra=False):
    p = ancora.insert_paragraph_before()
    p.style = ancora.part.document.styles["Normal"]
    fmt = p.paragraph_format
    fmt.alignment = alinhamento
    fmt.first_line_indent = Cm(0)
    fmt.space_before = Pt(antes)
    fmt.space_after = Pt(0)
    if entrelinhas == 1:
        fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    else:
        fmt.line_spacing = entrelinhas
    if recuo_esq is not None:
        fmt.left_indent = Cm(recuo_esq)
    partes = texto if isinstance(texto, list) else [(texto, negrito)]
    for trecho, neg in partes:
        r = p.add_run(trecho)
        r.bold = neg
        r.font.size = Pt(tamanho)
    if quebra:
        p.add_run().add_break(WD_BREAK.PAGE)
    return p


def secao_depois(par, corpo_sectpr, inicio=None):
    """Fecha uma seção no parágrafo `par`, com as mesmas dimensões do corpo."""
    s = copy.deepcopy(corpo_sectpr)
    for ref in s.findall(qn("w:headerReference")) + s.findall(qn("w:footerReference")):
        s.remove(ref)
    for pg in s.findall(qn("w:pgNumType")):
        s.remove(pg)
    if inicio is not None:
        pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), str(inicio)); s.append(pg)
    tipo = s.find(qn("w:type"))
    if tipo is None:
        tipo = OxmlElement("w:type"); s.insert(0, tipo)
    tipo.set(qn("w:val"), "nextPage")
    par._p.get_or_add_pPr().append(s)


def numero_de_pagina(header):
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Cm(0)
    campo(p, "PAGE", "1")
    for r in p._p.findall(qn("w:r")):  # fonte 10 no número da página (guia, p. 5)
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            r.insert(0, rpr)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "20")
        rpr.append(sz)


def main():
    faltando = [a for a in ARQUIVOS if not (CAPITULOS / a).exists()]
    presentes = [CAPITULOS / a for a in ARQUIVOS if (CAPITULOS / a).exists()]
    if faltando:
        print("aviso: capítulos ainda não escritos:", ", ".join(faltando), file=sys.stderr)
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    pandoc(presentes, SAIDA)

    ident = json.loads((AQUI / "identificacao.json").read_text(encoding="utf-8"))
    doc = Document(SAIDA)
    corpo_sectpr = doc.sections[-1]._sectPr

    # campos, espaço no número e títulos sem número
    primeiro_capitulo = None
    for p in doc.paragraphs:
        texto = p.text.strip()
        if texto.startswith("@@CAMPO:") and texto.endswith("@@"):
            chave = texto[len("@@CAMPO:"):-2]
            campo(p, CAMPOS[chave])
            continue
        nome = p.style.name if p.style is not None else ""
        if nome.startswith("Heading"):
            if p._p.find(".//" + qn("w:tab")) is not None:
                espaco_no_numero(p)
                if nome == "Heading 1" and primeiro_capitulo is None:
                    primeiro_capitulo = p
            elif nome == "Heading 1":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if nome == "Table Caption" and texto.startswith("Quadro "):
            p.style = doc.styles["Legenda de quadro"]

    # capa, folha de rosto e ficha (guia, p. 12-18), antes do primeiro parágrafo
    ancora = doc.paragraphs[0]
    titulo = [(ident["titulo"] + (":" if ident.get("subtitulo") else ""), True)]
    if ident.get("subtitulo"):
        titulo.append((" " + ident["subtitulo"], False))
    novo_paragrafo(ancora, ident["instituicao"])
    novo_paragrafo(ancora, ident["autor"])
    novo_paragrafo(ancora, titulo, antes=220)
    novo_paragrafo(ancora, ident["cidade"], antes=300)
    fim_capa = novo_paragrafo(ancora, ident["ano"])
    novo_paragrafo(ancora, ident["autor"])
    novo_paragrafo(ancora, titulo, antes=150)
    novo_paragrafo(ancora, ident["natureza"], alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=80,
                   entrelinhas=1, recuo_esq=8)
    novo_paragrafo(ancora, ident["cidade"], antes=240)
    novo_paragrafo(ancora, ident["ano"], quebra=True)
    novo_paragrafo(ancora, ident["ficha"], alinhamento=WD_ALIGN_PARAGRAPH.LEFT, antes=420, entrelinhas=1,
                   tamanho=10)
    # a ficha fica sozinha na página: o título seguinte (pré-textual) já quebra a página

    # seções: capa | pré-textuais (contadas a partir da folha de rosto) | texto (número impresso)
    secao_depois(fim_capa, corpo_sectpr)
    if primeiro_capitulo is not None:
        anterior = primeiro_capitulo._p.getprevious()
        while anterior is not None and anterior.tag != qn("w:p"):
            anterior = anterior.getprevious()
        if anterior is not None:
            from docx.text.paragraph import Paragraph
            secao_depois(Paragraph(anterior, primeiro_capitulo._parent), corpo_sectpr, inicio=1)
            primeiro_capitulo.paragraph_format.page_break_before = False
    secoes = doc.sections
    for s in secoes[:-1]:
        s.header.is_linked_to_previous = False
        for p in s.header.paragraphs:
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
    ultima = secoes[-1]
    ultima.header.is_linked_to_previous = False
    numero_de_pagina(ultima.header)
    for pg in ultima._sectPr.findall(qn("w:pgNumType")):
        ultima._sectPr.remove(pg)

    doc.settings.element.append(OxmlElement("w:updateFields"))
    doc.settings.element.find(qn("w:updateFields")).set(qn("w:val"), "true")
    doc.save(SAIDA)
    print(f"gravado: {SAIDA} ({len(presentes)} de {len(ARQUIVOS)} arquivos)")


if __name__ == "__main__":
    main()
